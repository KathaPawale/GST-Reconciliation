#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, date
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

MISSING_MARKERS = {"", "na", "n/a", "none", "null", "nan", "-"}

CANONICAL_COLUMNS = ("gstin", "invoice_number", "invoice_date", "igst", "cgst", "sgst")
COLUMN_ALIASES = {
    "gstin": {"gstin", "gstin no", "gstin number", "supplier gstin", "vendor gstin"},
    "invoice_number": {"invoice number", "invoice no", "invoice #", "bill number", "bill no", "bill #", "inv no"},
    "invoice_date": {"invoice date", "bill date", "inv date", "date"},
    "igst": {"igst", "igst amount", "i gst", "tax igst"},
    "cgst": {"cgst", "cgst amount", "c gst", "tax cgst"},
    "sgst": {"sgst", "sgst amount", "s gst", "tax sgst"},
}

YELLOW_FILL = PatternFill(start_color="FFF59D", end_color="FFF59D", fill_type="solid")
GREEN_FILL = PatternFill(start_color="C8E6C9", end_color="C8E6C9", fill_type="solid")
RED_FILL = PatternFill(start_color="FFCDD2", end_color="FFCDD2", fill_type="solid")


@dataclass
class ReconciledRow:
    source: str
    row_index: int
    original: Dict[str, Any]
    cleaned: Dict[str, Any]
    status: str
    reason: str
    match_key: Tuple[Optional[str], Optional[str], Optional[date]]
    duplicate_count: int = 1


@dataclass
class ReconciliationResult:
    b2b_rows: List[ReconciledRow]
    zoho_rows: List[ReconciledRow]
    only_b2b: List[ReconciledRow]
    only_zoho: List[ReconciledRow]


def normalize_header(name: str) -> str:
    return " ".join(str(name).strip().lower().replace("_", " ").split())


def normalize_missing(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    if normalize_header(text) in MISSING_MARKERS:
        return None
    return text


def parse_date(value: Any) -> Optional[date]:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = normalize_missing(value)
    if text is None:
        return None

    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%Y/%m/%d", "%m/%d/%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def parse_amount(value: Any) -> Optional[float]:
    text = normalize_missing(value)
    if text is None:
        return None
    normalized = text.replace(",", "")
    try:
        return float(normalized)
    except ValueError:
        return None


def canonical_column_map(headers: Iterable[str], mapping: Optional[Dict[str, str]] = None) -> Dict[str, str]:
    mapping = mapping or {}
    normalized_headers = {normalize_header(h): h for h in headers}
    result: Dict[str, str] = {}

    for canonical in CANONICAL_COLUMNS:
        if canonical in mapping:
            result[canonical] = mapping[canonical]
            continue

        for alias in COLUMN_ALIASES[canonical]:
            if alias in normalized_headers:
                result[canonical] = normalized_headers[alias]
                break

    missing = [key for key in CANONICAL_COLUMNS if key not in result]
    if missing:
        raise ValueError(
            "Missing required columns: "
            + ", ".join(missing)
            + ". Provide explicit mapping for these columns."
        )

    return result


def clean_record(raw: Dict[str, Any], column_map: Dict[str, str]) -> Dict[str, Any]:
    gstin = normalize_missing(raw.get(column_map["gstin"]))
    invoice = normalize_missing(raw.get(column_map["invoice_number"]))
    record_date = parse_date(raw.get(column_map["invoice_date"]))
    igst = parse_amount(raw.get(column_map["igst"]))
    cgst = parse_amount(raw.get(column_map["cgst"]))
    sgst = parse_amount(raw.get(column_map["sgst"]))

    return {
        "gstin": gstin.upper() if gstin else None,
        "invoice_number": invoice,
        "invoice_date": record_date,
        "igst": igst,
        "cgst": cgst,
        "sgst": sgst,
    }


def build_key(cleaned: Dict[str, Any]) -> Tuple[Optional[str], Optional[str], Optional[date]]:
    return cleaned["gstin"], cleaned["invoice_number"], cleaned["invoice_date"]


def all_tax_blank(cleaned: Dict[str, Any]) -> bool:
    return cleaned["igst"] is None and cleaned["cgst"] is None and cleaned["sgst"] is None


def taxes_match(b2b: Dict[str, Any], zoho: Dict[str, Any], tolerance: float) -> bool:
    for tax_key in ("igst", "cgst", "sgst"):
        left, right = b2b[tax_key], zoho[tax_key]
        if left is None and right is None:
            continue
        if left is None or right is None:
            return False
        if abs(left - right) > tolerance:
            return False
    return True


def infer_mismatch_reason(
    key: Tuple[Optional[str], Optional[str], Optional[date]],
    zoho_rows: List[Dict[str, Any]],
) -> str:
    gstin, inv, inv_date = key
    if inv is not None and inv_date is not None:
        for row in zoho_rows:
            if row["invoice_number"] == inv and row["invoice_date"] == inv_date and row["gstin"] != gstin:
                return "GSTIN mismatch"
    if gstin is not None and inv is not None:
        for row in zoho_rows:
            if row["gstin"] == gstin and row["invoice_number"] == inv and row["invoice_date"] != inv_date:
                return "Date mismatch"
    if gstin is not None and inv_date is not None:
        for row in zoho_rows:
            if row["gstin"] == gstin and row["invoice_date"] == inv_date and row["invoice_number"] != inv:
                return "Invoice/Bill Number mismatch"
    return "Record not found in counterpart sheet"


def reconcile_records(
    b2b_records: List[Dict[str, Any]],
    zoho_records: List[Dict[str, Any]],
    tolerance: float = 1.0,
) -> ReconciliationResult:
    b2b_keys = [build_key(row) for row in b2b_records]
    zoho_keys = [build_key(row) for row in zoho_records]
    b2b_dupes = Counter(b2b_keys)
    zoho_dupes = Counter(zoho_keys)

    zoho_by_key: Dict[Tuple[Optional[str], Optional[str], Optional[date]], List[int]] = defaultdict(list)
    for i, record in enumerate(zoho_records):
        zoho_by_key[build_key(record)].append(i)

    latest_b2b_date = max((r["invoice_date"] for r in b2b_records if r["invoice_date"]), default=None)
    prev_month = None
    if latest_b2b_date:
        if latest_b2b_date.month == 1:
            prev_month = (latest_b2b_date.year - 1, 12)
        else:
            prev_month = (latest_b2b_date.year, latest_b2b_date.month - 1)

    used_zoho_indexes: set[int] = set()
    b2b_result: List[ReconciledRow] = []
    zoho_result: List[ReconciledRow] = []
    only_b2b: List[ReconciledRow] = []

    for idx, b2b in enumerate(b2b_records):
        key = build_key(b2b)
        matched_idx = None
        for candidate_idx in zoho_by_key.get(key, []):
            if candidate_idx not in used_zoho_indexes:
                matched_idx = candidate_idx
                break

        if matched_idx is None:
            status = "MISMATCH"
            reason = infer_mismatch_reason(key, zoho_records)
            if prev_month and b2b["invoice_date"] and (b2b["invoice_date"].year, b2b["invoice_date"].month) == prev_month:
                status = "PREVIOUS_MONTH"
                reason = "Previous month's invoice"

            row = ReconciledRow(
                source="B2B",
                row_index=idx,
                original={},
                cleaned=b2b,
                status=status,
                reason=reason,
                match_key=key,
                duplicate_count=b2b_dupes[key],
            )
            b2b_result.append(row)
            only_b2b.append(row)
            continue

        used_zoho_indexes.add(matched_idx)
        zoho = zoho_records[matched_idx]

        if all_tax_blank(zoho):
            status = "ZOHO_TAX_BLANK"
            reason = "Zoho tax blank"
        elif taxes_match(b2b, zoho, tolerance):
            status = "MATCH"
            reason = "All tax amounts matched"
        else:
            status = "MISMATCH"
            reason = "Tax mismatch"

        b2b_row = ReconciledRow(
            source="B2B",
            row_index=idx,
            original={},
            cleaned=b2b,
            status=status,
            reason=reason,
            match_key=key,
            duplicate_count=b2b_dupes[key],
        )
        z_row = ReconciledRow(
            source="Zoho",
            row_index=matched_idx,
            original={},
            cleaned=zoho,
            status=status,
            reason=reason,
            match_key=key,
            duplicate_count=zoho_dupes[key],
        )
        b2b_result.append(b2b_row)
        zoho_result.append(z_row)

    only_zoho: List[ReconciledRow] = []
    for idx, zoho in enumerate(zoho_records):
        key = build_key(zoho)
        if idx not in used_zoho_indexes:
            row = ReconciledRow(
                source="Zoho",
                row_index=idx,
                original={},
                cleaned=zoho,
                status="MISMATCH",
                reason="Record not found in counterpart sheet",
                match_key=key,
                duplicate_count=zoho_dupes[key],
            )
            zoho_result.append(row)
            only_zoho.append(row)

    # Guarantee status/reason on every output row and preserve original ordering for zoho sheet
    zoho_result.sort(key=lambda r: r.row_index)
    return ReconciliationResult(b2b_result, zoho_result, only_b2b, only_zoho)


def read_worksheet_rows(path: Path, sheet_name: str) -> List[Dict[str, Any]]:
    wb = load_workbook(path, data_only=True)
    resolved_sheet = resolve_sheet_name(wb.sheetnames, sheet_name)
    if resolved_sheet is None:
        raise ValueError(f"Sheet '{sheet_name}' not found in {path}")
    ws = wb[resolved_sheet]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    data_rows: List[Dict[str, Any]] = []
    for row in rows[1:]:
        data_rows.append({headers[i]: row[i] if i < len(row) else None for i in range(len(headers))})
    return data_rows


def resolve_sheet_name(sheet_names: List[str], requested: str) -> Optional[str]:
    if requested in sheet_names:
        return requested

    def normalize(name: str) -> str:
        return "".join(ch.lower() for ch in name if ch.isalnum())

    requested_normalized = normalize(requested)
    for existing in sheet_names:
        if normalize(existing) == requested_normalized:
            return existing

    requested_without_slash = requested.replace("/", " ")
    requested_without_dash = requested.replace("/", "-")
    for candidate in (requested_without_slash, requested_without_dash):
        if candidate in sheet_names:
            return candidate

    if requested == "GST Portal B2B" and len(sheet_names) >= 1:
        return sheet_names[0]
    if requested == "Zoho Purchase/Bill" and len(sheet_names) >= 2:
        return sheet_names[1]

    return None


def read_csv_rows(path: Path) -> List[Dict[str, Any]]:
    with path.open(newline="", encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh)
        return list(reader)


def load_sources(
    input_file: Optional[Path],
    b2b_csv: Optional[Path],
    zoho_csv: Optional[Path],
    b2b_sheet: str,
    zoho_sheet: str,
) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    if input_file:
        b2b_raw = read_worksheet_rows(input_file, b2b_sheet)
        zoho_raw = read_worksheet_rows(input_file, zoho_sheet)
        return b2b_raw, zoho_raw
    if b2b_csv and zoho_csv:
        return read_csv_rows(b2b_csv), read_csv_rows(zoho_csv)
    raise ValueError("Provide either --input workbook.xlsx or both --b2b-csv and --zoho-csv")


def summarize_rows(rows: List[ReconciledRow]) -> Dict[str, Any]:
    def total(field: str) -> float:
        return sum(r.cleaned[field] or 0.0 for r in rows)

    counts = Counter(r.status for r in rows)
    return {
        "records": len(rows),
        "igst": total("igst"),
        "cgst": total("cgst"),
        "sgst": total("sgst"),
        "total_tax": total("igst") + total("cgst") + total("sgst"),
        "variance": 0.0,
        "yellow": counts.get("MATCH", 0),
        "green": counts.get("PREVIOUS_MONTH", 0),
        "red": counts.get("ZOHO_TAX_BLANK", 0),
        "grand_total": len(rows),
    }


def write_summary(ws, rows: List[ReconciledRow], start_col: int = 1) -> None:
    summary = summarize_rows(rows)
    labels = [
        ("Number of records", summary["records"]),
        ("Total IGST", summary["igst"]),
        ("Total CGST", summary["cgst"]),
        ("Total SGST", summary["sgst"]),
        ("Total tax", summary["total_tax"]),
        ("Total variance", summary["variance"]),
        ("Yellow count", summary["yellow"]),
        ("Green count", summary["green"]),
        ("Red count", summary["red"]),
        ("Grand total", summary["grand_total"]),
    ]
    ws.cell(row=1, column=start_col, value="Summary").font = Font(bold=True)
    for i, (label, value) in enumerate(labels, start=2):
        ws.cell(row=i, column=start_col, value=label)
        ws.cell(row=i, column=start_col + 1, value=value)


def write_result_sheet(ws, rows: List[ReconciledRow], title: str) -> None:
    write_summary(ws, rows)

    headers = [
        "GSTIN",
        "Invoice/Bill Number",
        "Invoice/Bill Date",
        "IGST",
        "CGST",
        "SGST",
        "Status",
        "Reason",
        "Duplicate Count",
        "Match Key",
    ]
    data_start = 13
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=data_start, column=col, value=header)
        cell.font = Font(bold=True)

    for row_num, row in enumerate(rows, start=data_start + 1):
        c = row.cleaned
        ws.cell(row=row_num, column=1, value=c["gstin"])
        ws.cell(row=row_num, column=2, value=c["invoice_number"])
        ws.cell(row=row_num, column=3, value=c["invoice_date"].isoformat() if c["invoice_date"] else None)
        ws.cell(row=row_num, column=4, value=c["igst"])
        ws.cell(row=row_num, column=5, value=c["cgst"])
        ws.cell(row=row_num, column=6, value=c["sgst"])
        ws.cell(row=row_num, column=7, value=row.status)
        ws.cell(row=row_num, column=8, value=row.reason)
        ws.cell(row=row_num, column=9, value=row.duplicate_count)
        ws.cell(row=row_num, column=10, value="|".join([str(v) if v is not None else "" for v in row.match_key]))

        if row.status == "MATCH":
            for tax_col in (4, 5, 6):
                ws.cell(row=row_num, column=tax_col).fill = YELLOW_FILL
        elif row.status == "PREVIOUS_MONTH" and title == "B2B":
            for col_idx in range(1, 11):
                ws.cell(row=row_num, column=col_idx).fill = GREEN_FILL
        elif row.status == "ZOHO_TAX_BLANK":
            for col_idx in range(1, 11):
                ws.cell(row=row_num, column=col_idx).fill = RED_FILL


def export_workbook(result: ReconciliationResult, output_path: Path) -> None:
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    sheets = [
        ("B2B", result.b2b_rows),
        ("Zoho", result.zoho_rows),
        ("Available in B2B but Not in Zoho", result.only_b2b),
        ("Available in Zoho but Not in B2B", result.only_zoho),
    ]

    for title, rows in sheets:
        ws = wb.create_sheet(title=title)
        write_result_sheet(ws, rows, title)

    wb.save(output_path)


def parse_mapping(raw: Optional[str]) -> Optional[Dict[str, Dict[str, str]]]:
    if not raw:
        return None
    parsed = json.loads(raw)
    if not isinstance(parsed, dict):
        raise ValueError("mapping-json must be a JSON object")
    return parsed


def run_reconciliation(
    input_file: Optional[Path],
    b2b_csv: Optional[Path],
    zoho_csv: Optional[Path],
    output_file: Path,
    tolerance: float,
    b2b_sheet: str,
    zoho_sheet: str,
    mapping_json: Optional[str] = None,
) -> ReconciliationResult:
    b2b_raw, zoho_raw = load_sources(input_file, b2b_csv, zoho_csv, b2b_sheet, zoho_sheet)

    if not b2b_raw or not zoho_raw:
        raise ValueError("Both B2B and Zoho source data must contain at least one row")

    mapping = parse_mapping(mapping_json) or {}

    b2b_map = canonical_column_map(b2b_raw[0].keys(), mapping.get("b2b"))
    zoho_map = canonical_column_map(zoho_raw[0].keys(), mapping.get("zoho"))

    b2b_clean = [clean_record(row, b2b_map) for row in b2b_raw]
    zoho_clean = [clean_record(row, zoho_map) for row in zoho_raw]

    result = reconcile_records(b2b_clean, zoho_clean, tolerance=tolerance)
    export_workbook(result, output_file)
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description="GST Portal vs Zoho reconciliation")
    parser.add_argument("--input", type=Path, help="Excel file containing B2B and Zoho sheets")
    parser.add_argument("--b2b-csv", type=Path, help="B2B CSV file")
    parser.add_argument("--zoho-csv", type=Path, help="Zoho CSV file")
    parser.add_argument("--output", type=Path, required=True, help="Output workbook path")
    parser.add_argument("--tolerance", type=float, default=1.0, help="Tax comparison tolerance")
    parser.add_argument("--b2b-sheet", default="GST Portal B2B")
    parser.add_argument("--zoho-sheet", default="Zoho Purchase/Bill")
    parser.add_argument(
        "--mapping-json",
        help='Optional JSON mapping, e.g. {"b2b":{"invoice_number":"Inv No"},"zoho":{...}}',
    )

    args = parser.parse_args()
    run_reconciliation(
        input_file=args.input,
        b2b_csv=args.b2b_csv,
        zoho_csv=args.zoho_csv,
        output_file=args.output,
        tolerance=args.tolerance,
        b2b_sheet=args.b2b_sheet,
        zoho_sheet=args.zoho_sheet,
        mapping_json=args.mapping_json,
    )


if __name__ == "__main__":
    main()
